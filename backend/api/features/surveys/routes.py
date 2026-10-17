from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from io import BytesIO
from sqlalchemy.orm import Session
from backend.api.core.database import get_db
from backend.api.features.surveys.models import SurveyModel
from backend.api.features.surveys.schemas import (
    SurveyOut, SurveyCreate, SurveyResponseSubmission, SurveySubmissionReceipt,
    SurveyStatusUpdate, SurveyAnalyticsOut, ReceiptVerificationOut, SurveyResponseStatusOut
)
from backend.api.features.surveys.service import (
    list_surveys, get_survey_by_id, create_survey, submit_survey_response,
    update_survey_status, get_survey_analytics, verify_survey_receipt, get_survey_user_response
)
from backend.api.middleware.auth_middleware import (
    get_current_user_token,
    get_optional_current_user_token,
)

router = APIRouter(prefix="/surveys", tags=["Surveys"])


def _survey_export_records(db: Session, survey_ids: list[str], owner_id: str | None = None):
    import json
    from backend.api.features.surveys.models import SurveyResponseModel

    surveys = db.query(SurveyModel).filter(SurveyModel.Id.in_(survey_ids)).all()
    if owner_id is not None:
        surveys = [survey for survey in surveys if survey.OwnerId == owner_id]
    responses = db.query(SurveyResponseModel).filter(
        SurveyResponseModel.SurveyId.in_([survey.Id for survey in surveys])
    ).order_by(SurveyResponseModel.Timestamp.desc()).all() if surveys else []
    response_map = {}
    for response in responses:
        response_map.setdefault(response.SurveyId, []).append({
            "timestamp": response.Timestamp,
            "respondent": response.VoterPseudonym,
            "answers": json.loads(response.AnswersJson or "[]"),
        })
    return surveys, response_map


def _build_export(surveys, response_map, export_format: str):
    if export_format == "csv":
        import csv
        from io import StringIO
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(["Survey ID", "Survey Title", "Response Timestamp", "Respondent", "Question", "Answer"])
        for survey in surveys:
            responses = response_map.get(survey.Id, [])
            if not responses:
                writer.writerow([survey.Id, survey.Title, "No responses recorded", "", "", ""])
            else:
                for resp in responses:
                    ts = resp.get("timestamp", "")
                    respondent = resp.get("respondent", "")
                    answers = resp.get("answers", [])
                    if not answers:
                        writer.writerow([survey.Id, survey.Title, ts, respondent, "None", ""])
                    else:
                        for ans in answers:
                            if isinstance(ans, dict):
                                q = ans.get("questionId") or ans.get("question") or "Question"
                                a = ans.get("answer", "")
                                writer.writerow([survey.Id, survey.Title, ts, respondent, q, a])
                            else:
                                writer.writerow([survey.Id, survey.Title, ts, respondent, "Answer", str(ans)])
        return output.getvalue().encode("utf-8"), "text/csv", "getitright-survey-responses.csv"

    if export_format == "json":
        import json
        data = []
        for survey in surveys:
            responses = response_map.get(survey.Id, [])
            data.append({
                "surveyId": survey.Id,
                "title": survey.Title,
                "status": survey.Status,
                "responseCount": len(responses),
                "responses": responses
            })
        return json.dumps(data, indent=2).encode("utf-8"), "application/json", "getitright-survey-responses.json"

    if export_format == "pdf":
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        story = [Paragraph("GetItRight survey responses", styles["Title"])]
        for survey in surveys:
            responses = response_map.get(survey.Id, [])
            story.extend([Spacer(1, 12), Paragraph(survey.Title, styles["Heading2"]), Paragraph(f"Responses: {len(responses)}", styles["BodyText"])])
            for index, response in enumerate(responses, 1):
                answers = "<br/>".join(
                    f"{answer.get('questionId', 'Question')}: {answer.get('answer', '')}"
                    for answer in response["answers"] if isinstance(answer, dict)
                )
                story.extend([Spacer(1, 8), Paragraph(f"Response {index} — {response['timestamp']}", styles["Heading3"]), Paragraph(answers or "No answers recorded.", styles["BodyText"])])
        document.build(story)
        return output.getvalue(), "application/pdf", "getitright-survey-responses.pdf"

    from docx import Document
    document = Document()
    document.add_heading("GetItRight survey responses", 0)
    for survey in surveys:
        responses = response_map.get(survey.Id, [])
        document.add_heading(survey.Title, level=1)
        document.add_paragraph(f"Responses: {len(responses)}")
        for index, response in enumerate(responses, 1):
            document.add_heading(f"Response {index} — {response['timestamp']}", level=2)
            document.add_paragraph(f"Respondent: {response['respondent']}")
            for answer in response["answers"]:
                if isinstance(answer, dict):
                    document.add_paragraph(f"{answer.get('questionId', 'Question')}: {answer.get('answer', '')}", style="List Bullet")
    output = BytesIO()
    document.save(output)
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "getitright-survey-responses.docx"

@router.get("", response_model=List[SurveyOut])
def get_all_surveys(
    plan: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    include_all: Optional[bool] = Query(False),
    db: Session = Depends(get_db)
):
    """Fetch surveys optionally filtered by plan and status."""
    return list_surveys(db, plan=plan, status=status, include_all=include_all)

@router.get("/templates")
def survey_templates_from_surveys():
    """Returns pre-built survey templates for the creation wizard (alias route)."""
    from backend.api.features.analytics.routes import survey_templates
    return survey_templates()


@router.get("/exports/responses")
def export_all_survey_responses(
    export_format: str = Query("pdf", alias="format", pattern="^(pdf|docx|csv|json)$"),
    owner_id: Optional[str] = Query(None),
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    if current_user["role"] != "Admin":
        owner_id = current_user["sub"]
    query = db.query(SurveyModel)
    if owner_id:
        query = query.filter(SurveyModel.OwnerId == owner_id)
    surveys = query.all()
    surveys, response_map = _survey_export_records(db, [survey.Id for survey in surveys], owner_id=owner_id)
    data, media_type, filename = _build_export(surveys, response_map, export_format)
    return StreamingResponse(BytesIO(data), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/{survey_id}/export")
def export_survey_responses(
    survey_id: str,
    export_format: str = Query("pdf", alias="format", pattern="^(pdf|docx|csv|json)$"),
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    survey = db.query(SurveyModel).filter(SurveyModel.Id == survey_id).first()
    if not survey:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Survey not found")
    if current_user["role"] != "Admin" and survey.OwnerId != current_user["sub"]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You do not own this survey")
    surveys, response_map = _survey_export_records(db, [survey_id])
    data, media_type, filename = _build_export(surveys, response_map, export_format)
    return StreamingResponse(BytesIO(data), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'})

@router.get("/{survey_id}/responses")
def get_survey_responses(
    survey_id: str,
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    """Return survey interactions to the owner or an administrator."""
    from backend.api.features.surveys.models import SurveyResponseModel
    import json

    survey = db.query(SurveyModel).filter(SurveyModel.Id == survey_id).first()
    if not survey:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Survey not found")
    is_admin = current_user.get("role") in ["Admin", "Super Admin"]
    if not is_admin and (survey.OwnerId or "").lower() != (current_user.get("sub") or "").lower():
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You do not own this survey")
    responses = db.query(SurveyResponseModel).filter(
        SurveyResponseModel.SurveyId == survey_id
    ).order_by(SurveyResponseModel.Timestamp.desc()).all()
    return [{
        "id": response.Id,
        "surveyId": response.SurveyId,
        "userId": response.UserId,
        "respondent": response.VoterPseudonym,
        "answers": json.loads(response.AnswersJson or "[]") if isinstance(response.AnswersJson, str) else (response.AnswersJson or []),
        "receiptHash": response.ReceiptHash,
        "timestamp": response.Timestamp,
        "voterName": getattr(response, "VoterName", None),
        "voterIdNumber": getattr(response, "VoterIdNumber", None),
        "deviceName": getattr(response, "DeviceName", None),
        "deviceOS": getattr(response, "DeviceOS", None),
        "county": getattr(response, "County", None),
        "latitude": getattr(response, "Latitude", None),
        "longitude": getattr(response, "Longitude", None),
        "gpsCoordinates": getattr(response, "GpsCoordinates", None),
    } for response in responses]

@router.get("/{survey_id}", response_model=SurveyOut)
def get_survey(survey_id: str, db: Session = Depends(get_db)):
    """Get single survey with full question and response breakdown."""
    return get_survey_by_id(db, survey_id)

@router.patch("/{survey_id}/status", response_model=SurveyOut)
def change_survey_status(
    survey_id: str,
    data: SurveyStatusUpdate,
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    """Update survey operational status (Active, Paused, Closed)."""
    survey = get_survey_by_id(db, survey_id)
    if current_user["role"] != "Admin" and survey.ownerId != current_user["sub"]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You do not own this survey")
    return update_survey_status(db, survey_id, data.status)

@router.get("/{survey_id}/analytics", response_model=SurveyAnalyticsOut)
def fetch_survey_analytics(
    survey_id: str,
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    """Fetch detailed question breakdown, completion rates, and Merkle root."""
    survey = get_survey_by_id(db, survey_id)
    if current_user["role"] != "Admin" and survey.ownerId != current_user["sub"]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You do not own this survey")
    return get_survey_analytics(db, survey_id)

@router.get("/{survey_id}/verify-receipt/{receipt_hash}", response_model=ReceiptVerificationOut)
def verify_receipt_endpoint(survey_id: str, receipt_hash: str, db: Session = Depends(get_db)):
    """Cryptographically verify survey respondent ballot receipt against Merkle tree."""
    return verify_survey_receipt(db, survey_id, receipt_hash)

@router.post("", response_model=SurveyOut)
def create_new_survey(
    data: SurveyCreate,
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    """Author and publish a new multi-step qualitative survey."""
    data.ownerId = current_user["sub"]
    data.plan = current_user["plan"]
    return create_survey(db, data)

@router.get("/{survey_id}/response-status", response_model=SurveyResponseStatusOut)
@router.get("/{survey_id}/status-check", response_model=SurveyResponseStatusOut)
def check_survey_user_status(
    survey_id: str,
    userId: Optional[str] = Query(None),
    voterKey: Optional[str] = Query(None),
    current_user: Optional[dict] = Depends(get_optional_current_user_token),
    db: Session = Depends(get_db)
):
    """Check if the user/voter has already taken this survey and return previous action receipt."""
    import json
    resp = get_survey_user_response(db, survey_id, user_id=current_user["sub"] if current_user else None, voter_key=voterKey)
    if resp:
        try:
            answers = json.loads(resp.AnswersJson)
        except Exception:
            answers = []
        return SurveyResponseStatusOut(
            hasResponded=True,
            surveyId=survey_id,
            receiptHash=resp.ReceiptHash,
            merkleLeafIndex=resp.MerkleLeafIndex,
            timestamp=resp.Timestamp,
            answers=answers,
            userId=resp.UserId,
            voterPseudonym=resp.VoterPseudonym
        )
    return SurveyResponseStatusOut(
        hasResponded=False,
        surveyId=survey_id
    )

@router.post("/{survey_id}/respond", response_model=SurveySubmissionReceipt)
@router.post("/{survey_id}/submit", response_model=SurveySubmissionReceipt)
def answer_survey(
    survey_id: str,
    sub: SurveyResponseSubmission,
    current_user: Optional[dict] = Depends(get_optional_current_user_token),
    db: Session = Depends(get_db),
):
    """Submit responses to a survey with Merkle receipt generation."""
    if current_user and not sub.userId:
        sub.userId = current_user["sub"]
    if not sub.userId:
        loc_tag = (sub.county or 'LOCATION1').replace(' ', '').upper()
        sub.userId = f"USR-{loc_tag}"
    if not sub.voterPseudonym:
        sub.voterPseudonym = sub.userId
    return submit_survey_response(db, survey_id, sub)
