from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from io import BytesIO
from sqlalchemy.orm import Session
from backend.api.core.database import get_db
from backend.api.features.polls.schemas import (
    PollOut, PollCreate, VoteCastRequest, VoteCastResponse, VoteTransactionOut, VoteStatusOut
)
from backend.api.features.polls.service import (
    list_polls, get_poll_by_id, create_poll, cast_vote, get_transactions_by_poll, reset_poll, get_vote_status
)

from backend.api.controllers.polls_controller import PollsController
from backend.api.middleware.auth_middleware import (
    get_current_user_token,
    get_optional_current_user_token,
)

router = APIRouter(prefix="/polls", tags=["Polls"])


def _build_poll_export(db: Session, owner_id: Optional[str], export_format: str):
    from backend.api.features.polls.models import PollModel, VoteTransactionModel

    query = db.query(PollModel)
    if owner_id:
        query = query.filter(PollModel.OwnerId == owner_id)
    polls = query.all()
    poll_ids = [poll.Id for poll in polls]
    transactions = db.query(VoteTransactionModel).filter(
        VoteTransactionModel.PollId.in_(poll_ids)
    ).order_by(VoteTransactionModel.Timestamp.desc()).all() if poll_ids else []
    transactions_by_poll = {}
    for transaction in transactions:
        transactions_by_poll.setdefault(transaction.PollId, []).append(transaction)

    if export_format == "csv":
        import csv
        from io import StringIO
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(["Poll ID", "Poll Title", "Category", "Status", "Vote Option", "Timestamp", "Voter Key", "Receipt Hash"])
        for poll in polls:
            votes = transactions_by_poll.get(poll.Id, [])
            if not votes:
                writer.writerow([poll.Id, poll.Title, poll.Category or "General", poll.Status, "No votes yet", "", "", ""])
            else:
                for v in votes:
                    writer.writerow([poll.Id, poll.Title, poll.Category or "General", poll.Status, v.OptionLabel, v.Timestamp, v.VoterPseudonym, v.ReceiptHash])
        return output.getvalue().encode("utf-8"), "text/csv", "getitright-poll-results.csv"

    if export_format == "json":
        import json
        data = []
        for poll in polls:
            votes = transactions_by_poll.get(poll.Id, [])
            data.append({
                "pollId": poll.Id,
                "title": poll.Title,
                "category": poll.Category,
                "status": poll.Status,
                "totalVotes": len(votes),
                "votes": [{
                    "id": v.Id,
                    "option": v.OptionLabel,
                    "timestamp": v.Timestamp,
                    "voter": v.VoterPseudonym,
                    "receiptHash": v.ReceiptHash
                } for v in votes]
            })
        return json.dumps(data, indent=2).encode("utf-8"), "application/json", "getitright-poll-results.json"

    if export_format == "pdf":
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        story = [Paragraph("GetItRight poll results", styles["Title"])]
        for poll in polls:
            votes = transactions_by_poll.get(poll.Id, [])
            story.extend([Spacer(1, 12), Paragraph(poll.Title, styles["Heading2"]), Paragraph(f"Votes: {len(votes)}", styles["BodyText"])])
            for index, vote in enumerate(votes, 1):
                story.append(Paragraph(f"Vote {index} — {vote.Timestamp}: {vote.OptionLabel}", styles["BodyText"]))
        document.build(story)
        return output.getvalue(), "application/pdf", "getitright-poll-results.pdf"

    from docx import Document
    document = Document()
    document.add_heading("GetItRight poll results", 0)
    for poll in polls:
        votes = transactions_by_poll.get(poll.Id, [])
        document.add_heading(poll.Title, level=1)
        document.add_paragraph(f"Votes: {len(votes)}")
        for index, vote in enumerate(votes, 1):
            document.add_paragraph(f"Vote {index} — {vote.Timestamp}: {vote.OptionLabel}", style="List Bullet")
    output = BytesIO()
    document.save(output)
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "getitright-poll-results.docx"

@router.get("", response_model=List[PollOut])
def get_polls(
    plan: Optional[str] = Query(None, description="Filter polls by tier plan"),
    status: Optional[str] = Query(None, description="Filter polls by status"),
    include_all: Optional[bool] = Query(False, description="Include archived/disabled polls (admin)"),
    db: Session = Depends(get_db)
):
    """Fetch real polls directly from database, optionally filtered by plan and status."""
    return PollsController.get_polls(db, plan=plan, status=status, include_all=include_all)


@router.get("/exports/results")
def export_poll_results(
    export_format: str = Query("pdf", alias="format", pattern="^(pdf|docx|csv|json)$"),
    owner_id: Optional[str] = Query(None),
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    if current_user["role"] != "Admin":
        owner_id = current_user["sub"]
    data, media_type, filename = _build_poll_export(db, owner_id, export_format)
    return StreamingResponse(BytesIO(data), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/{poll_id}/export")
def export_single_poll(
    poll_id: str,
    export_format: str = Query("csv", alias="format", pattern="^(pdf|docx|csv|json)$"),
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    from backend.api.features.polls.models import PollModel, VoteTransactionModel
    poll = db.query(PollModel).filter(PollModel.Id == poll_id).first()
    if not poll:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Poll not found")
    if current_user["role"] != "Admin" and poll.OwnerId != current_user["sub"]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You do not own this poll")

    votes = db.query(VoteTransactionModel).filter(
        VoteTransactionModel.PollId == poll_id
    ).order_by(VoteTransactionModel.Timestamp.desc()).all()

    slug = "".join(c for c in poll.Title.lower().replace(" ", "-") if c.isalnum() or c == "-")[:30] or "poll"

    if export_format == "csv":
        import csv
        from io import StringIO
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(["Poll ID", "Poll Title", "Category", "Status", "Vote Option", "Timestamp", "Voter Key", "Receipt Hash"])
        if not votes:
            writer.writerow([poll.Id, poll.Title, poll.Category or "General", poll.Status, "No votes yet", "", "", ""])
        else:
            for v in votes:
                writer.writerow([poll.Id, poll.Title, poll.Category or "General", poll.Status, v.OptionLabel, v.Timestamp, v.VoterPseudonym, v.ReceiptHash])
        return StreamingResponse(BytesIO(output.getvalue().encode("utf-8")), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{slug}-votes.csv"'})

    if export_format == "json":
        import json
        data = {
            "pollId": poll.Id,
            "title": poll.Title,
            "category": poll.Category,
            "status": poll.Status,
            "totalVotes": len(votes),
            "votes": [{
                "id": v.Id,
                "option": v.OptionLabel,
                "timestamp": v.Timestamp,
                "voter": v.VoterPseudonym,
                "receiptHash": v.ReceiptHash
            } for v in votes]
        }
        return StreamingResponse(BytesIO(json.dumps(data, indent=2).encode("utf-8")), media_type="application/json", headers={"Content-Disposition": f'attachment; filename="{slug}-votes.json"'})

    if export_format == "pdf":
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        story = [
            Paragraph(f"Poll Results: {poll.Title}", styles["Title"]),
            Spacer(1, 10),
            Paragraph(f"Category: {poll.Category or 'General'} | Status: {poll.Status} | Total Votes: {len(votes)}", styles["Normal"]),
            Spacer(1, 14),
            Paragraph("Vote Transactions Log", styles["Heading2"])
        ]
        for idx, v in enumerate(votes, 1):
            story.append(Paragraph(f"{idx}. {v.Timestamp} — <strong>{v.OptionLabel}</strong> (Receipt: {v.ReceiptHash[:16]}...)", styles["BodyText"]))
        document.build(story)
        return StreamingResponse(BytesIO(output.getvalue()), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{slug}-results.pdf"'})

    from docx import Document
    document = Document()
    document.add_heading(poll.Title, 0)
    document.add_paragraph(f"Category: {poll.Category or 'General'} | Status: {poll.Status} | Total Votes: {len(votes)}")
    document.add_heading("Vote Transactions Log", level=1)
    for idx, v in enumerate(votes, 1):
        document.add_paragraph(f"{idx}. {v.Timestamp}: {v.OptionLabel} (Voter: {v.VoterPseudonym})", style="List Bullet")
    output = BytesIO()
    document.save(output)
    return StreamingResponse(BytesIO(output.getvalue()), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{slug}-results.docx"'})

@router.get("/{poll_id}", response_model=PollOut)
def get_single_poll(poll_id: str, db: Session = Depends(get_db)):
    """Fetch a single poll by ID with real option counts and calculated percentages."""
    return PollsController.get_single_poll(db, poll_id)

@router.get("/{poll_id}/vote-status", response_model=VoteStatusOut)
def check_vote_status(
    poll_id: str,
    userId: Optional[str] = Query(None),
    voterKey: Optional[str] = Query(None),
    current_user: Optional[dict] = Depends(get_optional_current_user_token),
    db: Session = Depends(get_db)
):
    """Check if a user or voter pseudonym has already voted on this poll."""
    trusted_user_id = current_user["sub"] if current_user else None
    return get_vote_status(db, poll_id, user_id=trusted_user_id, voter_key=voterKey)

@router.post("", response_model=PollOut)
def create_new_poll(
    data: PollCreate,
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    """Create a new poll with custom voting options."""
    data.ownerId = current_user["sub"]
    data.plan = current_user["plan"]
    return PollsController.create_poll(db, data)

@router.post("/{poll_id}/vote", response_model=VoteCastResponse)
def submit_vote(
    poll_id: str,
    vote_req: VoteCastRequest,
    current_user: Optional[dict] = Depends(get_optional_current_user_token),
    db: Session = Depends(get_db),
):
    """Cast a verifiable ballot, recalculate percentage distribution, and issue a cryptographic receipt."""
    # Retain optional user ID for audit tracing if provided, but do not force-lockout session votes
    if current_user and not vote_req.userId:
        vote_req.userId = current_user.get("sub")
    loc_tag = (vote_req.county or "LOCATION1").replace(" ", "").upper()
    if not (vote_req.voterName and vote_req.voterName.strip()):
        vote_req.voterName = f"USER@{loc_tag}"
    if not (vote_req.voterKey and vote_req.voterKey.strip()):
        import uuid
        vote_req.voterKey = f"USER@{loc_tag}-{uuid.uuid4().hex[:6].upper()}"
    return PollsController.submit_vote(db, poll_id, vote_req)

@router.post("/{poll_id}/reset", response_model=PollOut)
def reset_poll_endpoint(poll_id: str):
    """Retired: rewriting a tally breaks its verifiable record."""
    raise HTTPException(status_code=status.HTTP_410_GONE, detail="Poll resets are disabled. Publish a correction poll instead.")

@router.get("/{poll_id}/transactions", response_model=List[VoteTransactionOut])
def get_poll_audit_transactions(
    poll_id: str,
    current_user: dict = Depends(get_current_user_token),
    db: Session = Depends(get_db),
):
    """Retrieve audit transactions and Merkle receipts for a specific poll."""
    poll = get_poll_by_id(db, poll_id)
    is_admin = current_user.get("role") in ["Admin", "Super Admin"]
    if not is_admin and (poll.ownerId or "").lower() != (current_user.get("sub") or "").lower():
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You do not own this poll")
    return get_transactions_by_poll(db, poll_id)
